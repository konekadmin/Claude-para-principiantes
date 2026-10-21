"""Genera los datos ficticios del video.

Salidas:
  material/gastos_sep2025_vs_sep2026.xlsx  -> el "export del ERP" que se adjunta a Claude
  produccion/escenas/data.js               -> los mismos datos para los dashboards y las escenas

Todos los montos son ficticios, en soles (S/).
"""

import json
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent

AREAS = ["Ventas", "Operaciones", "TI", "RR.HH.", "Finanzas", "Marketing"]
CATEGORIAS = ["Personal", "Servicios", "Viajes", "Software", "Alquiler"]
CENTROS = {
    "Ventas": "CC-VTA-100",
    "Operaciones": "CC-OPE-200",
    "TI": "CC-TIC-300",
    "RR.HH.": "CC-RHH-400",
    "Finanzas": "CC-FIN-500",
    "Marketing": "CC-MKT-600",
}

# Gasto de septiembre por área y categoría: (2025, 2026)
GASTO = {
    "Ventas": {
        "Personal": (520_340, 548_120),
        "Servicios": (84_910, 82_350),
        "Viajes": (63_880, 49_160),
        "Software": (38_250, 41_030),
        "Alquiler": (71_620, 71_340),
    },
    "Operaciones": {
        "Personal": (689_730, 741_860),
        "Servicios": (210_480, 236_270),
        "Viajes": (21_940, 24_310),
        "Software": (26_110, 27_040),
        "Alquiler": (179_740, 267_520),
    },
    "TI": {
        "Personal": (310_160, 326_480),
        "Servicios": (119_870, 130_940),
        "Viajes": (8_120, 6_870),
        "Software": (144_960, 231_650),
        "Alquiler": (29_890, 30_060),
    },
    "RR.HH.": {
        "Personal": (180_420, 176_350),
        "Servicios": (95_160, 70_880),
        "Viajes": (9_040, 6_120),
        "Software": (17_930, 24_270),
        "Alquiler": (25_450, 25_380),
    },
    "Finanzas": {
        "Personal": (259_870, 268_140),
        "Servicios": (70_230, 64_410),
        "Viajes": (6_010, 4_870),
        "Software": (31_860, 35_120),
        "Alquiler": (28_030, 28_460),
    },
    "Marketing": {
        "Personal": (209_640, 214_380),
        "Servicios": (239_870, 297_960),
        "Viajes": (30_150, 36_240),
        "Software": (22_080, 29_310),
        "Alquiler": (20_310, 20_110),
    },
}


def filas():
    for anio, idx in ((2025, 0), (2026, 1)):
        for area in AREAS:
            for cat in CATEGORIAS:
                yield {
                    "anio": anio,
                    "mes": "Septiembre",
                    "periodo": f"{anio}-09",
                    "area": area,
                    "categoria": cat,
                    "centro": CENTROS[area],
                    "importe": GASTO[area][cat][idx],
                }


def resumen():
    tot = {2025: 0, 2026: 0}
    por_area = {}
    por_cat = {c: [0, 0] for c in CATEGORIAS}
    for area in AREAS:
        a = [0, 0]
        for cat in CATEGORIAS:
            v25, v26 = GASTO[area][cat]
            a[0] += v25
            a[1] += v26
            por_cat[cat][0] += v25
            por_cat[cat][1] += v26
        por_area[area] = a
        tot[2025] += a[0]
        tot[2026] += a[1]
    return tot, por_area, por_cat


def escribir_excel(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Gastos"
    cab = ["Año", "Mes", "Periodo", "Área", "Categoría", "Centro de costo", "Moneda", "Importe"]
    ws.append(cab)
    for f in filas():
        ws.append([f["anio"], f["mes"], f["periodo"], f["area"], f["categoria"], f["centro"], "PEN", f["importe"]])

    thin = Side(style="thin", color="BFBFBF")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="595959")
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        for c in row:
            c.number_format = "#,##0.00"
    anchos = [7, 12, 10, 14, 12, 16, 9, 14]
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{ws.max_row}"

    nota = wb.create_sheet("Nota")
    nota["A1"] = "Datos ficticios para práctica. No corresponden a ninguna empresa real."
    nota["A2"] = "Fuente simulada: exportación del ERP, gasto de septiembre por área y categoría, 2025 vs 2026, en soles (PEN)."
    nota.column_dimensions["A"].width = 110
    wb.save(path)
    fijar_fechas(path)


def fijar_fechas(path, fecha="2026-10-18T00:00:00Z"):
    """openpyxl siempre graba la fecha actual; se fija para que el Excel sea idéntico en cada ejecución."""
    with zipfile.ZipFile(path) as z:
        partes = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, datos in partes:
            if info.filename == "docProps/core.xml":
                txt = datos.decode("utf-8")
                for etiqueta in ("dcterms:created", "dcterms:modified"):
                    a = txt.index(f"<{etiqueta}")
                    a = txt.index(">", a) + 1
                    b = txt.index(f"</{etiqueta}>", a)
                    txt = txt[:a] + fecha + txt[b:]
                datos = txt.encode("utf-8")
            info.date_time = datetime(2026, 10, 18).timetuple()[:6]
            z.writestr(info, datos)


def escribir_js(path: Path):
    tot, por_area, por_cat = resumen()
    data = {
        "areas": AREAS,
        "categorias": CATEGORIAS,
        "gasto": {a: {c: list(GASTO[a][c]) for c in CATEGORIAS} for a in AREAS},
        "total": [tot[2025], tot[2026]],
        "porArea": por_area,
        "porCategoria": por_cat,
        "filas": list(filas()),
    }
    path.write_text(
        "// Generado por produccion/datos.py — datos ficticios (S/)\n"
        f"window.KONEK_DATA = {json.dumps(data, ensure_ascii=False)};\n",
        encoding="utf-8",
    )


def main():
    tot, por_area, por_cat = resumen()
    escribir_excel(BASE / "material" / "gastos_sep2025_vs_sep2026.xlsx")
    escribir_js(BASE / "produccion" / "escenas" / "data.js")

    var = tot[2026] - tot[2025]
    print(f"Total 2025: S/ {tot[2025]:,}  2026: S/ {tot[2026]:,}  var: S/ {var:,} ({var / tot[2025]:.1%})")
    for a, (v25, v26) in por_area.items():
        print(f"  {a:12} {v25:>10,} {v26:>10,} {v26 - v25:>+9,} {(v26 - v25) / v25:+.1%}  peso var {(v26 - v25) / var:.0%}")
    for c, (v25, v26) in por_cat.items():
        print(f"  {c:12} {v25:>10,} {v26:>10,} {v26 - v25:>+9,} {(v26 - v25) / v25:+.1%}  peso 2026 {v26 / tot[2026]:.1%}")


if __name__ == "__main__":
    main()
